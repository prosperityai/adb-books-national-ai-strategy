# Purpose: tests for finance/build_budget.py (formulas present, totals consistent).
# Book: National AI Strategy in Practice, Chapter 10
# Licence: MIT (see LICENSE)
import copy

import build_budget as bb
import pytest
from openpyxl import load_workbook


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    inputs = bb.load_inputs()
    items = bb.all_line_items(inputs)
    path = tmp_path_factory.mktemp("budget") / "budget.xlsx"
    bb.build_workbook(inputs, items).save(path)
    return inputs, items, load_workbook(path)


def test_sheets_and_formulas_present(built):
    _, _, wb = built
    assert wb.sheetnames == ["Read me", "Line items", "Financing shares", "By pillar",
                             "By source", "Summary"]
    assert wb["By pillar"]["B2"].value.startswith("=SUMIFS(")
    assert wb["By source"]["B2"].value.startswith("=SUMPRODUCT(")
    assert wb["Summary"]["B2"].value.startswith("=")
    # No typed totals: every cell right of the labels on the totals sheets is a formula.
    for name in ("By pillar", "By source"):
        for row in wb[name].iter_rows(min_row=2, min_col=2):
            assert all(str(c.value).startswith("=") for c in row)
    line_items = wb["Line items"]
    assert all(str(line_items.cell(row=r, column=9).value).startswith("=SUM(")
               for r in range(2, line_items.max_row + 1))


def test_compute_lines_come_from_planner(built):
    _, items, _ = built
    compute = [row for row in items if row[2] == "From compute planner"]
    assert len(compute) == 2 and sum(compute[0][3:]) > 0


def test_totals_balance(built):
    inputs, items, _ = built
    by_pillar, by_source = bb.totals(inputs, items)
    grand = sum(sum(row[3:]) for row in items)
    assert sum(sum(v) for v in by_pillar.values()) == pytest.approx(grand)
    assert sum(sum(v) for v in by_source.values()) == pytest.approx(grand)


def test_bad_financing_shares_rejected():
    inputs = copy.deepcopy(bb.load_inputs())
    inputs["financing_shares"]["Data"] = [0.5, 0.0, 0.0, 0.0, 0.0]
    with pytest.raises(ValueError):
        bb.validate_shares(inputs)
